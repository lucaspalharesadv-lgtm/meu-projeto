#!/usr/bin/env python3
"""Excel dos 200 contatos de correspondencia: abas Contatos e Textos dos e-mails."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
rows = list(csv.DictReader(open(f"{HERE}/correspondencia/lista_200.csv", encoding="utf8")))
NAVY = "14284B"; hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor=NAVY)
wb = Workbook(); ws = wb.active; ws.title = "Contatos"
ws.append(["Nº", "Categoria", "Empresa / escritório", "E-mail", "Site", "Status"])
for r in rows: ws.append([int(r["n"]), r["categoria"], r["empresa"], r["email"], r["site"], r["status"]])
ws2 = wb.create_sheet("Textos dos e-mails"); ws2.append(["Nº", "E-mail (destinatário)", "Assunto", "Texto do e-mail", "Status"])
for r in rows: ws2.append([int(r["n"]), r["email"], r["assunto"], r["corpo"], r["status"]])
for sh, widths in ((ws, [5, 28, 36, 42, 44, 12]), (ws2, [5, 42, 52, 105, 12])):
    for c in range(1, len(widths) + 1):
        h = sh.cell(1, c); h.font = hf; h.fill = hfill; h.alignment = Alignment(horizontal="center", vertical="center")
        sh.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in sh.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    sh.freeze_panes = "B2"; sh.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{sh.max_row}"
for row in range(2, ws2.max_row + 1): ws2.row_dimensions[row].height = 200
out = f"{HERE}/Correspondencia_200_Contatos_Lucas_Palhares.xlsx"; wb.save(out); print("ok", len(rows), "contatos ->", out)
