#!/usr/bin/env python3
"""Gera Correspondencia_Contatos_Lucas_Palhares.xlsx (e-mails, links de cadastro, sem canal, texto de oferta)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__)); D = f"{HERE}/correspondencia"
rows = []
for f, tag in (("plataformas.tsv", "Plataforma / gestora"), ("escritorios.tsv", "Escritório / depto. jurídico")):
    for l in open(f"{D}/{f}", encoding="utf8"):
        c = (l.rstrip("\n").split("\t") + [""] * 8)[:8]
        if c[0] and c[0] != "ORGANIZACAO": rows.append([tag] + c)
em = [r for r in rows if "@" in r[4]]
fm = [r for r in rows if "@" not in r[4] and r[6].startswith("http")]
so = [r for r in rows if "@" not in r[4] and not r[6].startswith("http")]
NAVY = "14284B"; hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor=NAVY)
wb = Workbook()
def sheet(name, header, data, widths, first=False):
    ws = wb.active if first else wb.create_sheet(name); ws.title = name; ws.append(header)
    for r in data: ws.append(r)
    for c in range(1, len(header) + 1):
        h = ws.cell(1, c); h.font = hf; h.fill = hfill; h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"; return ws
sheet("E-mails", ["Nº", "Empresa", "Tipo", "E-mail", "Site", "Observação", "Status"],
      [[i, r[1], r[0], r[4], r[3], r[8], ""] for i, r in enumerate(em, 1)], [5, 36, 26, 38, 38, 60, 12], first=True)
sheet("Cadastros (formulário)", ["Nº", "Empresa", "Tipo", "Link de cadastro / correspondente", "Site", "Observação", "Status"],
      [[i, r[1], r[0], r[6], r[3], r[8], ""] for i, r in enumerate(fm, 1)], [5, 40, 26, 62, 38, 50, 12])
sheet("Sem canal (revisar)", ["Nº", "Empresa", "Tipo", "Site", "Telefone / WhatsApp", "Observação"],
      [[i, r[1], r[0], r[3], r[5], r[8]] for i, r in enumerate(so, 1)], [5, 44, 26, 44, 22, 50])
txt = """Assunto: Correspondente jurídico em Ji-Paraná/RO: audiências e diligências

Prezados,

Sou advogado (OAB/RO 11.037), de Ji-Paraná/RO, e gostaria de me cadastrar como correspondente jurídico para audiências, diligências, cópias processuais e protocolos na minha região.

Tenho mais de 5 anos de contencioso, com audiências, sustentações orais e acompanhamento de processos no PJe, e-SAJ, eproc e Projudi. Estagiei na Justiça Federal e na AGU/Procuradoria Federal em Ji-Paraná, e por isso conheço bem a rotina desses órgãos por aqui. Atuo com pontualidade e comunicação clara.

Peço que me informem como funciona o cadastro, os valores por serviço e o prazo de pagamento. Meu currículo segue em anexo.

Se preferirem não receber novas mensagens, basta responder e não voltarei a escrever.

Atenciosamente,
Lucas Alexandre Horas Palhares
Advogado | OAB/RO 11.037
(69) 99335-9788 | lucaspalharesadv@gmail.com"""
ws = wb.create_sheet("Texto do e-mail"); ws.column_dimensions["A"].width = 110
for i, line in enumerate(txt.split("\n"), 1): ws.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(bold=True, color=NAVY)
out = f"{HERE}/Correspondencia_Contatos_Lucas_Palhares.xlsx"; wb.save(out); print("ok", len(em), "e-mails |", len(fm), "cadastros |", len(so), "sem canal")
