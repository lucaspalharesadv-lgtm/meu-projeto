#!/usr/bin/env python3
"""Gera emails/Contatos_Orgaos_Publicos_Lucas_Palhares.xlsx a partir de publico/resultado_publico.csv e emails_prontos.csv."""
import csv, re, collections, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
res = list(csv.DictReader(open(f"{HERE}/publico/resultado_publico.csv", encoding="utf8")))
pro = {x["email"]: x for x in csv.DictReader(open(f"{HERE}/publico/emails_prontos.csv", encoding="utf8"))}
NOME = {"judiciario_federal": "Judiciário federal e superior", "judiciario_estadual": "Judiciário estadual",
        "mp_defensorias": "Ministérios Públicos e Defensorias", "procuradorias_controle": "Procuradorias e controle",
        "legislativo_federal_estadual": "Legislativo federal e estadual", "camaras_municipais": "Câmaras municipais",
        "executivo_autarquias": "Executivo, autarquias e universidades", "estatais_sistemas_conselhos": "Estatais, Sistema S e conselhos",
        "ji_parana_regiao": "Ji-Paraná e região (RO)"}
ORDEM = list(NOME)
res.sort(key=lambda x: (ORDEM.index(x["categoria"]), x["uf"] or "zz", x["orgao"], x["email"]))
NAVY = "14284B"
hf = Font(bold=True, color="FFFFFF", name="Calibri", size=11); hfill = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="D5DBE6"); border = Border(top=thin, bottom=thin, left=thin, right=thin)
org_curto = lambda o: re.split(r"\s[-–]\s|\(", o)[0].strip()

wb = Workbook()
# ---- Resumo ----
ws = wb.active; ws.title = "Resumo"
ws["A1"] = "Contatos institucionais de órgãos públicos"; ws["A1"].font = Font(bold=True, size=16, color=NAVY)
ws["A2"] = "Lucas Alexandre Horas Palhares  |  OAB/RO 11.037  |  coleta de 03/10/2026"; ws["A2"].font = Font(italic=True, color="555555")
ws["A4"] = "Categoria"; ws["B4"] = "Contatos"
for c in ("A4", "B4"): ws[c].font = hf; ws[c].fill = hfill; ws[c].alignment = Alignment(horizontal="center")
cnt = collections.Counter(x["categoria"] for x in res); r = 5
for c in ORDEM: ws.cell(r, 1, NOME[c]); ws.cell(r, 2, cnt[c]); r += 1
ws.cell(r, 1, "TOTAL").font = Font(bold=True); ws.cell(r, 2, f"=SUM(B5:B{r-1})").font = Font(bold=True); r += 2
ws.cell(r, 1, "Por tipo de caixa").font = Font(bold=True, color=NAVY); r += 1
for t, n in collections.Counter(x["tipo"] for x in res).most_common(): ws.cell(r, 1, t); ws.cell(r, 2, n); r += 1
r += 1; ws.cell(r, 1, "Por confiança").font = Font(bold=True, color=NAVY); r += 1
for t, n in collections.Counter(x["confianca"] for x in res).most_common():
    ws.cell(r, 1, {"alta": "Alta (confirmado na página oficial)", "media": "Média (só em resultado de busca; pode devolver erro)"}[t]); ws.cell(r, 2, n); r += 1
r += 1
for n in ['Como usar: a aba "Contatos" tem filtros em cada coluna; a aba "Textos dos e-mails" tem o assunto e o corpo de cada mensagem.',
          'A coluna Status controla o envio: vazio = não enviado; "rascunho" = rascunho criado no Gmail; "enviado"; "não contatar" = pediu para parar.',
          "São caixas institucionais de cargo ou setor, publicadas pelos próprios órgãos. Não há dados pessoais nem e-mails de webmail.",
          "Em Ji-Paraná só entram contatos gerais (RH, protocolo, secretaria), nunca gestores, e a frase de mudança não aparece nesses e-mails.",
          "Nenhum e-mail foi enviado ainda. Cada e-mail deve sair COM o currículo em PDF anexado."]:
    ws.cell(r, 1, n).alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 32; r += 1
ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 12
for c in "CDEF": ws.column_dimensions[c].width = 18

# ---- Contatos ----
ws2 = wb.create_sheet("Contatos")
cols = ["Nº", "Categoria", "Órgão", "Esfera", "UF", "Cidade", "Setor / cargo", "E-mail", "Tipo", "Confiança", "Ji-Paraná", "Fonte (página oficial)", "Status"]
ws2.append(cols)
for i, x in enumerate(res, 1):
    ws2.append([i, NOME[x["categoria"]], org_curto(x["orgao"]), x["esfera"].capitalize(), x["uf"] or "", x["cidade"], x["setor_cargo"], x["email"], x["tipo"],
                x["confianca"].capitalize(), "Sim" if x["ji_parana"] == "sim" else "Não", x["fonte"], pro[x["email"]]["status"] if x["email"] in pro else ""])
for c in range(1, len(cols) + 1):
    h = ws2.cell(1, c); h.font = hf; h.fill = hfill; h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws2.row_dimensions[1].height = 30
for i, w in enumerate([5, 34, 40, 11, 6, 20, 34, 42, 20, 11, 10, 50, 13], 1): ws2.column_dimensions[get_column_letter(i)].width = w
band = PatternFill("solid", fgColor="F2F5FA")
for row in range(2, ws2.max_row + 1):
    for c in range(1, len(cols) + 1):
        cell = ws2.cell(row, c); cell.border = border; cell.alignment = Alignment(vertical="top")
        if row % 2 == 0: cell.fill = band
    ws2.cell(row, 8).font = Font(color="1F3A68", bold=True)
    if ws2.cell(row, 10).value == "Média": ws2.cell(row, 10).font = Font(color="B45309")
ws2.freeze_panes = "C2"; ws2.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws2.max_row}"

# ---- Textos ----
ws3 = wb.create_sheet("Textos dos e-mails")
ws3.append(["Nº", "E-mail (destinatário)", "Órgão", "Assunto", "Texto do e-mail", "Status"])
for i, x in enumerate(res, 1):
    p = pro[x["email"]]; ws3.append([i, x["email"], org_curto(x["orgao"]), p["assunto"], p["corpo"], p["status"]])
for c in range(1, 7):
    h = ws3.cell(1, c); h.font = hf; h.fill = hfill; h.alignment = Alignment(horizontal="center", vertical="center")
for i, w in enumerate([5, 40, 34, 52, 110, 13], 1): ws3.column_dimensions[get_column_letter(i)].width = w
for row in range(2, ws3.max_row + 1):
    for c in range(1, 7): ws3.cell(row, c).alignment = Alignment(vertical="top", wrap_text=True)
    ws3.row_dimensions[row].height = 215
ws3.freeze_panes = "C2"; ws3.auto_filter.ref = f"A1:F{ws3.max_row}"
out = f"{HERE}/Contatos_Orgaos_Publicos_Lucas_Palhares.xlsx"
wb.save(out); print("ok", len(res), "contatos ->", out)
