# -*- coding: utf-8 -*-
"""Seleciona os proximos N contatos para rascunho, na ordem combinada, pulando quem ja esta em controle_envios.csv.
Uso: python3 proximos.py 40  -> imprime JSON com os contatos e TODAS as colunas de personalizacao."""
import csv, json, os, re, sys
from openpyxl import load_workbook
B = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CTRL = os.path.join(B, 'rotina', 'controle_envios.csv')
BLOQ = re.compile(r'\bmbt\b|mbtadvoca|mbtadvogados|ernesto ?borges|ernestoborges|pessoa ?(&|e) ?pessoa|pessoaepessoa', re.I)
PORD = {'Alta': 0, 'Media': 1, 'Média': 1, 'Baixa': 2}

def ler(xlsx, aba='Contatos'):
    ws = load_workbook(os.path.join(B, xlsx), read_only=True)[aba]
    rows = list(ws.iter_rows(values_only=True)); h = [str(x or '').strip() for x in rows[0]]
    return [dict(zip(h, [('' if v is None else str(v)) for v in r])) for r in rows[1:] if r and r[h.index('E-mail')]]

def ja_contatados():
    if not os.path.exists(CTRL): return set()
    return {r['E-mail'].strip().lower() for r in csv.DictReader(open(CTRL, encoding='utf-8-sig'))}

def fila():
    P = ler('PARCERIAS_ASSOCIADO_500.xlsx'); O = ler('PROSPECCAO_JURIDICA_BRASIL.xlsx')
    for r in P: r['_planilha'] = 'PARCERIAS_ASSOCIADO_500'
    for r in O: r['_planilha'] = 'PROSPECCAO_JURIDICA_BRASIL'; r['Lote'] = '0'
    def chave(r):
        lote = (r.get('Lote') or '0')[:1]
        ordem_lote = {'3': 0, '2': 1, '1': 2, '0': 3}.get(lote, 4)
        p = PORD.get(r.get('Prioridade', ''), 1)
        return (p if p == 2 else 0, ordem_lote, p)   # Baixas de todas as planilhas por ultimo
    return sorted(P + O, key=chave)

def proximos(n):
    feitos = ja_contatados(); out = []; vistos = set()
    for r in fila():
        e = r['E-mail'].strip().lower()
        if e in feitos or e in vistos or BLOQ.search(json.dumps(r, ensure_ascii=False)): continue
        vistos.add(e); out.append(r)
        if len(out) >= n: break
    return out

if __name__ == '__main__':
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 40
    sel = proximos(n)
    restam = len([1 for r in fila() if r['E-mail'].strip().lower() not in ja_contatados()])
    print(json.dumps({'selecionados': sel, 'restam_na_fila': restam}, ensure_ascii=False, indent=1))
